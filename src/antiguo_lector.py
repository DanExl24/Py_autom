import sys
import os

# Añadir el directorio raíz al path para permitir importar 'auth'
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from auth.servicio_general import obtener_servicio_drive
import io
import re
import json
import datetime
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook

drive = obtener_servicio_drive()
PASAN2026 = {}
hoja_min_row = 0
hoja_max_row = 0
hoja_min_col = 0
hoja_max_col = 0
CARPETA_PRINCIPAL = "1nc5qTNXcP05h4N8AMka-CODYrGTXk4rm"

# Rutas dinámicas basadas en la ubicación del script
base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
output_dir = os.path.join(base_dir, "output")
os.makedirs(output_dir, exist_ok=True)

path_redes = os.path.join(output_dir, "Red_conocimiento.json")
path_encabezados = os.path.join(output_dir, "Encabezados.json")
path_fichas = os.path.join(output_dir, "fichas.json")

# Intentar cargar variables desde el archivo existente o inicializarlas a cero
try:
    if os.path.exists(path_redes):
        with open(path_redes, "r", encoding="utf-8") as archivo:
            ars = json.load(archivo)
            ar = ars.get("FORMACIONES TITULADA REGULAR MODALIDAD PRESENCIAL Y VIRTUAL 2026")
            if ar:
                hoja_min_row = int(ar['min_row'])
                hoja_max_row = int(ar['max_row'])
                hoja_min_col = int(ar['min_col'])
                hoja_max_col = int(ar['max_col'])
except Exception as e:
    print(f"Advertencia al leer Red_conocimiento.json: {e}")

# ========== FUNCIONES HELPER DE LIMPIEZA ==========

def sanitizar_nombre(texto):
    """Limpia los nombres de encabezados y valores de texto:
    - Corrige tildes rotas por encoding (DISEÑO, GESTIÓN, etc.)
    - Elimina saltos de línea y espacios dobles
    - Corrige el cero por O en INICI0 → INICIO
    """
    if not isinstance(texto, str):
        return texto
    
    # Reemplazar saltos de línea por espacio
    texto = texto.replace('\n', ' ')
    
    # Eliminar espacios dobles o más
    while '  ' in texto:
        texto = texto.replace('  ', ' ')
    
    # Corregir el cero por O en INICI0
    texto = texto.replace('INICI0', 'INICIO')
    
    # Corregir tildes rotas por problemas de encoding comunes en Excel
    # Primero: reemplazar el caracter de reemplazo Unicode (aparece cuando openpyxl no puede decodificar)
    bad_char = "\ufffd"
    texto = texto.replace(f"DISE{bad_char}O", "DISEÑO")
    texto = texto.replace(f"GESTI{bad_char}N", "GESTIÓN")
    texto = texto.replace(f"CONSTRUCCI{bad_char}N", "CONSTRUCCIÓN")
    texto = texto.replace(f"ANIMACI{bad_char}N", "ANIMACIÓN")
    texto = texto.replace(f"TECN{bad_char}LOGO", "TECNÓLOGO")
    texto = texto.replace(f"T{bad_char}CNICO", "TÉCNICO")
    texto = texto.replace(f"C{bad_char}DIGO", "CÓDIGO")
    texto = texto.replace(f"A{bad_char}O", "AÑO")
    texto = texto.replace(f"MONTA{bad_char}I", "MONTAÑI")
    texto = texto.replace(f"UNI{bad_char}N", "UNIÓN")
    
    # Segundo: reemplazos de texto plano (sin tildes → con tildes)
    reemplazos = {
        "TECNOLOGO": "TECNÓLOGO",
        "TECNICO": "TÉCNICO",
    }
    for buscar, reemplazar in reemplazos.items():
        texto = texto.replace(buscar, reemplazar)
    
    return texto.strip()

def limpiar_valor(valor):
    """Limpia los tipos de datos que vienen del Excel:
    - None, "None", "NAN", "" → None (null en JSON)
    - datetime → string "YYYY-MM-DD"
    - float 23.0 → int 23
    - strings → sanitizados con tildes corregidas
    """
    if valor is None:
        return None
    
    # Strings: verificar si representan un valor nulo, luego sanitizar
    if isinstance(valor, str):
        valor_limpio = valor.strip()
        if valor_limpio.upper() in ('NONE', 'NAN', ''):
            return None
        return sanitizar_nombre(valor_limpio)
    
    # Fechas de Python/Excel → formato estándar "YYYY-MM-DD"
    if isinstance(valor, (datetime.datetime, datetime.date)):
        return valor.strftime('%Y-%m-%d')
    
    # Floats enteros → int limpio (23.0 → 23)
    if isinstance(valor, float):
        if valor.is_integer():
            return int(valor)
        return valor
    
    return valor

def normalizar_trimestre(valor):
    """Normaliza strings de trimestre a formato estándar T-[ROMANO]-[AÑO].
    Ejemplos: "T- II-2025" → "T-II-2025", "T-III 2025" → "T-III-2025"
    """
    if not isinstance(valor, str):
        return valor
    s = valor.strip().upper()
    match = re.search(r'T\s*-\s*(I{1,3}|IV|V)\s*[-/ ]\s*(\d{4})', s)
    if match:
        return f"T-{match.group(1)}-{match.group(2)}"
    return s

def resolver_celda(hoja, fila, columna):
    """Resuelve el valor real de una celda, incluso si está dentro
    de un rango combinado (merged cells). Si la celda pertenece a un
    rango combinado, retorna el valor de la celda origen (esquina superior izquierda).
    """
    for rango in hoja.merged_cells.ranges:
        if (rango.min_row <= fila <= rango.max_row and 
            rango.min_col <= columna <= rango.max_col):
            return hoja.cell(row=rango.min_row, column=rango.min_col).value
    return hoja.cell(row=fila, column=columna).value

# ========== TU LÓGICA ORIGINAL (SIN CAMBIOS ESTRUCTURALES) ==========

def carpeta_principal(ID):
    resultado = (
    drive.files()  # type: ignore
    .list(
        q=f"'{ID}' in parents",
        fields="files(id,name,mimeType,webViewLink)"
    )
    .execute()
    )
    return resultado

# Obtener archivos dentro de la carpeta principal
oferta = carpeta_principal(CARPETA_PRINCIPAL).get("files",[])
programas_regulares = {}
for archivo in oferta:
    if (archivo['name']=='CONSOLIDADO PROGRAMAS REGULAR - 2026.xlsx'):
        programas_regulares = archivo

# Hacer la peticion a google drive al archivo de programas regulares por medio de la ID del archivo
request = drive.files().get_media(fileId=programas_regulares['id'])  # type: ignore

# crear un archivo en memoria RAM
excel = io.BytesIO()

# Descargar el archivo de drive y escribir sus datos dentro de la variable excel, asi se genera el archivo en memoria
downloader = MediaIoBaseDownload(excel, request)

# Variable para saber cuando termina la descarga
done = False
# Descargar el archivo
while not done:
    _, done = downloader.next_chunk() #Descargar por pedazos (chuncks)

excel.seek(0) #Leer desde el comienzo del cursor

wb = load_workbook(excel, data_only=True)

# print(wb.sheetnames)

def leer_hoja ():
    hoja = wb["PASAN 2026 - TI-TII-TIII"]
    # print(f"Hoja: {hoja}")
    # print(f"Filas de la hoja: {hoja.max_row}")
    # print(f"Columnas de la hoja: {hoja.max_column}")
    # print(f"Dimensiones calculadas: {hoja.calculate_dimension()}")
    return hoja

def debug_celdas():
    hoja = leer_hoja()

    with open("celdas.txt", "w", encoding="utf-8") as archivo:
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.value is not None:
                    archivo.write(f"Coordenada: {celda.coordinate}, valor: {celda.value}\n")

# leer_hoja('PASAN 2026 ')

def obtener_rango(nombre):
    hoja = leer_hoja()
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value is not None:
                if(celda.value == nombre):
                    for rango in hoja.merged_cells.ranges:
                        if celda.coordinate in rango:
                            print(f"Rango: {rango}")
                            return {"Rango": {rango}}
                        else:
                            print(f"Celda: {celda.coordinate}")
                    return {"Celda" : {celda.coordinate}}
                
def obtener_rango_hoja(nombre):
    hoja = leer_hoja()
    rg = {}
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value is not None:
                if(celda.value == nombre):
                    for rango in hoja.merged_cells.ranges:
                        if celda.coordinate in rango:
                            PASAN2026[nombre] = {
                                "nombre" : nombre,
                                "coordenadas" : f"{rango}",
                                "min_row" : f"{rango.min_row}",
                                "max_row" : f"{rango.max_row}",
                                "min_col" : f"{rango.min_col}",
                                "max_col" : f"{rango.max_col}",
                            }
                            return PASAN2026
                        else:
                            print(f"No hay rango")
                        

#with open("Red_conocimiento.json", "w") as seed_redes:
#    json.dump(obtener_rango_hoja("FORMACIONES TITULADA REGULAR MODALIDAD PRESENCIAL Y VIRTUAL 2026"),seed_redes,ensure_ascii=False,indent=1)

# Funcion para obtener las redes de conocimiento
def obtener_redes():
    hoja = leer_hoja()
    #print(type(json))
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value is not None:
                if(celda.value == "RED DE CONOCIMIENTO"):
                    # Recorrer todas las redes sumando el valor de las filas
                    FilaRedes = celda.row+1
                    CeldaRedes = hoja.cell(FilaRedes,celda.column)
                    print(CeldaRedes.value)
                    encontrado = False
                    for rango in hoja.merged_cells.ranges:
                        if CeldaRedes.coordinate in rango:
                            print(rango)
                            PASAN2026[CeldaRedes.value] = {
                                "nombre" : CeldaRedes.value,
                                "coordenadas" : f"{rango}",
                                "min_row" : f"{rango.min_row}",
                                "max_row" : f"{rango.max_row}",
                                "min_col" : f"{rango.min_col}",
                                "max_col" : f"{rango.max_col}",
                            }
                            encontrado = True
                            break
                    if not encontrado and CeldaRedes.value is not None:
                        print(f"Red de fila única: {CeldaRedes.coordinate}")
                        PASAN2026[CeldaRedes.value] = {
                            "nombre" : CeldaRedes.value,
                            "coordenadas" : CeldaRedes.coordinate,
                            "min_row" : f"{CeldaRedes.row}",
                            "max_row" : f"{CeldaRedes.row}",
                            "min_col" : f"{CeldaRedes.column}",
                            "max_col" : f"{CeldaRedes.column}",
                        }
    return PASAN2026
                    #print(celda.row)
                    #print(celda.column)

# Las variables hoja_min_row, etc., se leen dinámicamente al inicio o durante la ejecución.


                        
def consultar_valor(rango):
    hoja = leer_hoja()
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value is not None:
                if celda.coordinate in rango:
                    print(f"Valor: {celda.value}")




def obtener_encabezados():
    hoja = leer_hoja()
    encabezado = {}
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value is not None:
                if(celda.value == "RED DE CONOCIMIENTO"):
                    print(celda.coordinate)
                    ColInicioFila = celda.row
                    ColInicio = celda.column
                    print("ColInicio:", ColInicio)
                    print("hoja_max_col:", hoja_max_col)
                    for rangoEncabezado in range(ColInicio, hoja_max_col + 1):
                        encabezados = hoja.cell(row=ColInicioFila,column=rangoEncabezado) 
                        encabezado[encabezados.value] = {
                            "nombre" : encabezados.value,
                            "coordenadas" : encabezados.coordinate,
                            "row" : encabezados.row,
                            "col" : encabezados.column
                        }
                    return encabezado
#with open("Encabezados.json", "w") as encabezados:
#    json.dump(obtener_encabezados(),encabezados,ensure_ascii=False,indent=1)



def ficha_por_red():
    hoja = leer_hoja()
    redes = {}
    fichas = {}
    encabezado = {}
    with open(path_redes, "r", encoding="utf-8") as redes_json:
        redes = json.load(redes_json)
    with open(path_encabezados, "r", encoding="utf-8") as encabezados:
        encabezado = json.load(encabezados)
    for nombre_red, valores_red in redes.items():
        # Ignorar la entrada de metadatos general del rango de la hoja
        if nombre_red == "FORMACIONES TITULADA REGULAR MODALIDAD PRESENCIAL Y VIRTUAL 2026":
            continue
            
        fichas_red = {}
        coordenadas = valores_red
        for filas in range(int(coordenadas["min_row"]),int(coordenadas["max_row"]) + 1):
            registro = {}
            for nombre_encabezado,valores_encabezado in encabezado.items():
                columna = valores_encabezado["col"]
                
                # MEJORA 1 + 4: Resolver celdas combinadas y limpiar el valor
                valor_crudo = resolver_celda(hoja, filas, columna)
                valor_limpio = limpiar_valor(valor_crudo)
                
                # Sanitizar el nombre del encabezado (quita espacios dobles, tildes rotas, etc.)
                nombre_limpio = sanitizar_nombre(nombre_encabezado)
                
                # MEJORA 2: Normalizar trimestres
                if "TRIMESTRE" in nombre_limpio.upper() and isinstance(valor_limpio, str):
                    valor_limpio = normalizar_trimestre(valor_limpio)
                
                registro[nombre_limpio] = valor_limpio
                
                if nombre_limpio == "RED DE CONOCIMIENTO":
                    registro["RED DE CONOCIMIENTO"] = sanitizar_nombre(nombre_red)
                    
            numero_ficha = registro.get("FICHA")
            
            # Ignorar filas vacías, encabezados repetidos o fichas no numéricas (ej. títulos)
            if not numero_ficha:
                continue
            ficha_str = str(numero_ficha).strip()
            if not ficha_str.isdigit() or ficha_str.upper() == "FICHA":
                continue
                
            fichas_red[ficha_str] = registro
        fichas[sanitizar_nombre(nombre_red)] = fichas_red
    return fichas

if __name__ == "__main__":
    # Siempre regenerar los JSONs auxiliares para evitar desalineación cuando cambie la hoja
    print("Generando Red_conocimiento.json...")
    obtener_rango_hoja("FORMACIONES TITULADA REGULAR MODALIDAD PRESENCIAL Y VIRTUAL 2026")
    obtener_redes()
    with open(path_redes, "w", encoding="utf-8") as f:
        json.dump(PASAN2026, f, ensure_ascii=False, indent=1)
        
    # Cargar las dimensiones de la hoja a partir de la nueva generación
    ar = PASAN2026.get("FORMACIONES TITULADA REGULAR MODALIDAD PRESENCIAL Y VIRTUAL 2026")
    if ar:
        hoja_min_row = int(ar['min_row'])
        hoja_max_row = int(ar['max_row'])
        hoja_min_col = int(ar['min_col'])
        hoja_max_col = int(ar['max_col'])
        
    print("Generando Encabezados.json...")
    enc_datos = obtener_encabezados()
    with open(path_encabezados, "w", encoding="utf-8") as f:
        json.dump(enc_datos, f, ensure_ascii=False, indent=1)
        
    # Procesar y guardar el archivo final fichas.json
    print("Extrayendo fichas y guardando en fichas.json...")
    fichas_finales = ficha_por_red()
    with open(path_fichas, "w", encoding="utf-8") as f:
        json.dump(fichas_finales, f, ensure_ascii=False, indent=1)
    print("Sincronización finalizada con éxito!")
